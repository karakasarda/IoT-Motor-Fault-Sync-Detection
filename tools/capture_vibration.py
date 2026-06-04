import argparse
import math
import re
import time
from pathlib import Path

import matplotlib.pyplot as plt
import openpyxl
import serial
from openpyxl.chart import LineChart, Reference

DEFAULT_OUTPUT_DIR = Path(__file__).resolve().parents[1] / "data" / "captures"
DEFAULT_OUTPUT_PREFIX = "vibration_capture"

LEGACY_FEATURE_HEADERS = [
    "time_ms",
    "ax",
    "ay",
    "az",
    "gx",
    "gy",
    "gz",
    "acc_mag",
    "gyro_mag",
    "label",
]
FEATURE_HEADERS = [
    "time_ms",
    "ax",
    "ay",
    "az",
    "gx",
    "gy",
    "gz",
    "acc_mag",
    "gyro_mag",
    "pulse_count",
    "edge_count",
    "rise_count",
    "fall_count",
    "state",
    "pulse_rate_hz",
    "last_edge_dt_ms",
    "label",
]
FEATURE_HEADER_LINE = ",".join(FEATURE_HEADERS)
LEGACY_FEATURE_HEADER_LINE = ",".join(LEGACY_FEATURE_HEADERS)
HEADER_LINES = {FEATURE_HEADER_LINE, LEGACY_FEATURE_HEADER_LINE}
LEGACY_LINE_RE = re.compile(
    r"accel\(g\)=\(([-0-9.]+), ([-0-9.]+), ([-0-9.]+)\) "
    r"gyro\(dps\)=\(([-0-9.]+), ([-0-9.]+), ([-0-9.]+)\) "
    r"temp=([-0-9.]+)C roll=([-0-9.]+) pitch=([-0-9.]+)"
)


def format_serial_error(port, exc):
    return (
        f"Serial port {port} acilamadi: {exc}\n"
        "COM port baska bir uygulama tarafindan kullaniyor olabilir.\n"
        "Kapatilacak olasi kaynaklar: MicroPico terminali, canli dashboard/grafik "
        "penceresi, baska seri port araci veya VS Code icindeki acik Pico baglantisi.\n"
        "Portlari kontrol etmek icin: python -m serial.tools.list_ports -v"
    )


def parse_args():
    parser = argparse.ArgumentParser(
        description="Capture live vibration data from a Pico-based IMU stream."
    )
    parser.add_argument("--port", default="COM10", help="Serial port of the Pico")
    parser.add_argument(
        "--baudrate", type=int, default=115200, help="Serial baudrate"
    )
    parser.add_argument(
        "--duration",
        type=float,
        default=10.0,
        help="Capture duration in seconds",
    )
    parser.add_argument(
        "--output-dir",
        default=str(DEFAULT_OUTPUT_DIR),
        help="Folder where CSV/XLSX/PNG files will be written",
    )
    parser.add_argument(
        "--soft-reset",
        action="store_true",
        help="Send Ctrl+D before capture to restart main.py cleanly",
    )
    parser.add_argument(
        "--label",
        default=None,
        help="Override the label column for all captured rows",
    )
    parser.add_argument(
        "--file-prefix",
        default=DEFAULT_OUTPUT_PREFIX,
        help="Prefix used for CSV/XLSX/PNG/LOG output files",
    )
    return parser.parse_args()


def parse_stream_line(line, fallback_time_ms=None, label_override=None):
    if not line or line.startswith("#") or line in HEADER_LINES:
        return None

    parts = [part.strip() for part in line.split(",")]
    if len(parts) == len(FEATURE_HEADERS):
        try:
            sample = {
                "time_ms": int(float(parts[0])),
                "ax": float(parts[1]),
                "ay": float(parts[2]),
                "az": float(parts[3]),
                "gx": float(parts[4]),
                "gy": float(parts[5]),
                "gz": float(parts[6]),
                "acc_mag": float(parts[7]),
                "gyro_mag": float(parts[8]),
                "pulse_count": int(float(parts[9])),
                "edge_count": int(float(parts[10])),
                "rise_count": int(float(parts[11])),
                "fall_count": int(float(parts[12])),
                "state": int(float(parts[13])),
                "pulse_rate_hz": float(parts[14]),
                "last_edge_dt_ms": int(float(parts[15])),
                "label": parts[16],
                "raw_line": line,
            }
            if label_override is not None:
                sample["label"] = label_override
            return sample
        except ValueError:
            pass

    if len(parts) == len(LEGACY_FEATURE_HEADERS):
        try:
            sample = {
                "time_ms": int(float(parts[0])),
                "ax": float(parts[1]),
                "ay": float(parts[2]),
                "az": float(parts[3]),
                "gx": float(parts[4]),
                "gy": float(parts[5]),
                "gz": float(parts[6]),
                "acc_mag": float(parts[7]),
                "gyro_mag": float(parts[8]),
                "pulse_count": 0,
                "edge_count": 0,
                "rise_count": 0,
                "fall_count": 0,
                "state": 0,
                "pulse_rate_hz": 0.0,
                "last_edge_dt_ms": -1,
                "label": parts[9],
                "raw_line": line,
            }
            if label_override is not None:
                sample["label"] = label_override
            return sample
        except ValueError:
            pass

    match = LEGACY_LINE_RE.search(line)
    if not match:
        return None

    ax = float(match.group(1))
    ay = float(match.group(2))
    az = float(match.group(3))
    gx = float(match.group(4))
    gy = float(match.group(5))
    gz = float(match.group(6))

    return {
        "time_ms": int(fallback_time_ms or 0),
        "ax": ax,
        "ay": ay,
        "az": az,
        "gx": gx,
        "gy": gy,
        "gz": gz,
        "acc_mag": math.sqrt(ax * ax + ay * ay + az * az),
        "gyro_mag": math.sqrt(gx * gx + gy * gy + gz * gz),
        "pulse_count": 0,
        "edge_count": 0,
        "rise_count": 0,
        "fall_count": 0,
        "state": 0,
        "pulse_rate_hz": 0.0,
        "last_edge_dt_ms": -1,
        "label": label_override or "legacy",
        "raw_line": line,
    }


def capture_samples(port, baudrate, duration, soft_reset, label_override=None):
    samples = []
    raw_lines = []

    with serial.Serial(port, baudrate, timeout=0.25, write_timeout=1) as ser:
        time.sleep(1.0)
        raw_lines.extend(prepare_stream(ser, soft_reset))

        start = time.time()
        buffer = ""

        while time.time() - start < duration:
            chunk = ser.read(4096).decode("utf-8", errors="replace")
            if not chunk:
                continue

            buffer += chunk

            while "\n" in buffer:
                line, buffer = buffer.split("\n", 1)
                line = line.strip()
                if not line:
                    continue
                raw_lines.append(line)

                sample = parse_stream_line(
                    line,
                    fallback_time_ms=round((time.time() - start) * 1000),
                    label_override=label_override,
                )
                if sample is not None:
                    samples.append(sample)

    return samples, raw_lines


def prepare_stream(ser, soft_reset):
    ser.reset_input_buffer()
    ser.reset_output_buffer()

    if not soft_reset:
        return []

    ser.write(b"\x03\x03")
    ser.flush()
    time.sleep(0.3)
    ser.reset_input_buffer()
    ser.reset_output_buffer()
    ser.write(b"\x04")
    ser.flush()
    return wait_for_stream_header(ser)


def wait_for_stream_header(ser, timeout=6.0):
    start = time.time()
    buffer = ""
    raw_lines = []

    while time.time() - start < timeout:
        chunk = ser.read(4096).decode("utf-8", errors="replace")
        if not chunk:
            continue

        buffer += chunk
        while "\n" in buffer:
            line, buffer = buffer.split("\n", 1)
            line = line.strip()
            if not line:
                continue
            raw_lines.append(line)
            if line in HEADER_LINES:
                return raw_lines

    return raw_lines


def ensure_output_paths(output_dir, file_prefix=DEFAULT_OUTPUT_PREFIX):
    folder = Path(output_dir)
    folder.mkdir(parents=True, exist_ok=True)
    stamp = time.strftime("%Y%m%d_%H%M%S")
    return {
        "folder": folder,
        "csv": folder / f"{file_prefix}_{stamp}.csv",
        "xlsx": folder / f"{file_prefix}_{stamp}.xlsx",
        "png": folder / f"{file_prefix}_{stamp}.png",
        "log": folder / f"{file_prefix}_{stamp}.log",
    }


def write_csv(samples, path):
    with path.open("w", encoding="utf-8", newline="") as f:
        f.write(FEATURE_HEADER_LINE + "\n")
        for sample in samples:
            f.write(",".join(str(sample[h]) for h in FEATURE_HEADERS) + "\n")


def write_log(raw_lines, path):
    path.write_text("\n".join(raw_lines), encoding="utf-8")


def create_plot(samples, path):
    t = [sample["time_ms"] / 1000.0 for sample in samples]

    fig, axes = plt.subplots(4, 1, figsize=(12, 12), sharex=True)

    axes[0].plot(t, [sample["ax"] for sample in samples], label="ax")
    axes[0].plot(t, [sample["ay"] for sample in samples], label="ay")
    axes[0].plot(t, [sample["az"] for sample in samples], label="az")
    axes[0].set_ylabel("Acceleration (g)")
    axes[0].set_title("MPU6050 Compatible IMU Capture")
    axes[0].grid(True, alpha=0.3)
    axes[0].legend(loc="upper right")

    axes[1].plot(t, [sample["gx"] for sample in samples], label="gx")
    axes[1].plot(t, [sample["gy"] for sample in samples], label="gy")
    axes[1].plot(t, [sample["gz"] for sample in samples], label="gz")
    axes[1].set_ylabel("Gyro (dps)")
    axes[1].grid(True, alpha=0.3)
    axes[1].legend(loc="upper right")

    axes[2].plot(t, [sample["acc_mag"] for sample in samples], label="acc_mag")
    axes[2].set_ylabel("Acc Magnitude")
    axes[2].grid(True, alpha=0.3)
    axes[2].legend(loc="upper right")

    axes[3].plot(t, [sample["gyro_mag"] for sample in samples], label="gyro_mag")
    axes[3].set_xlabel("Time (s)")
    axes[3].set_ylabel("Gyro Magnitude")
    axes[3].grid(True, alpha=0.3)
    axes[3].legend(loc="upper right")

    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def add_chart(ws, title, y_axis_title, min_col, max_col, max_row, anchor):
    chart = LineChart()
    chart.title = title
    chart.y_axis.title = y_axis_title
    chart.x_axis.title = "Time (ms)"
    chart.height = 8
    chart.width = 18

    data = Reference(ws, min_col=min_col, max_col=max_col, min_row=1, max_row=max_row)
    cats = Reference(ws, min_col=1, min_row=2, max_row=max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, anchor)


def build_summary(samples):
    if not samples:
        return {
            "sample_count": 0,
            "duration_ms": 0,
            "duration_s": 0.0,
            "sample_rate_hz": 0.0,
            "label": "",
            "pulse_count": 0,
            "pulse_rate_hz_avg": 0.0,
            "pulse_rate_hz_max": 0.0,
        }

    duration_ms = samples[-1]["time_ms"] - samples[0]["time_ms"]
    duration_s = duration_ms / 1000.0 if duration_ms else 0.0
    sample_rate_hz = (len(samples) / duration_s) if duration_s > 0 else 0.0
    pulse_rates = [float(sample.get("pulse_rate_hz", 0.0)) for sample in samples]
    return {
        "sample_count": len(samples),
        "duration_ms": duration_ms,
        "duration_s": duration_s,
        "sample_rate_hz": sample_rate_hz,
        "label": samples[-1]["label"],
        "pulse_count": int(samples[-1].get("pulse_count", 0)),
        "pulse_rate_hz_avg": sum(pulse_rates) / len(pulse_rates) if pulse_rates else 0.0,
        "pulse_rate_hz_max": max(pulse_rates) if pulse_rates else 0.0,
    }


def create_excel(samples, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "data"

    ws.append(FEATURE_HEADERS)

    for sample in samples:
        ws.append([sample[h] for h in FEATURE_HEADERS])

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 14

    summary = wb.create_sheet("summary")
    summary.append(["metric", "value"])
    capture_summary = build_summary(samples)
    summary.append(["sample_count", capture_summary["sample_count"]])
    summary.append(["duration_ms", capture_summary["duration_ms"]])
    summary.append(["duration_s", capture_summary["duration_s"]])
    summary.append(["sample_rate_hz", capture_summary["sample_rate_hz"]])
    summary.append(["label", capture_summary["label"]])
    summary.append(["pulse_count", capture_summary.get("pulse_count", 0)])
    summary.append(["pulse_rate_hz_avg", capture_summary.get("pulse_rate_hz_avg", 0.0)])
    summary.append(["pulse_rate_hz_max", capture_summary.get("pulse_rate_hz_max", 0.0)])

    metrics = [
        "ax",
        "ay",
        "az",
        "gx",
        "gy",
        "gz",
        "acc_mag",
        "gyro_mag",
        "pulse_rate_hz",
    ]
    for metric in metrics:
        values = [sample[metric] for sample in samples]
        summary.append([f"{metric}_min", min(values)])
        summary.append([f"{metric}_max", max(values)])
        summary.append([f"{metric}_avg", sum(values) / len(values)])

    max_row = len(samples) + 1
    add_chart(ws, "Acceleration", "g", 2, 4, max_row, "L2")
    add_chart(ws, "Gyroscope", "dps", 5, 7, max_row, "L20")
    add_chart(ws, "Magnitudes", "magnitude", 8, 9, max_row, "L38")

    wb.save(path)


def main():
    args = parse_args()
    paths = ensure_output_paths(args.output_dir, args.file_prefix)

    try:
        samples, raw_lines = capture_samples(
            args.port,
            args.baudrate,
            args.duration,
            args.soft_reset,
            label_override=args.label,
        )
    except serial.SerialException as exc:
        raise SystemExit(format_serial_error(args.port, exc)) from exc

    if not samples:
        raise SystemExit("No sensor samples were captured from the Pico stream.")

    write_csv(samples, paths["csv"])
    write_log(raw_lines, paths["log"])
    create_plot(samples, paths["png"])
    create_excel(samples, paths["xlsx"])

    summary = build_summary(samples)
    print(f"Captured {summary['sample_count']} samples.")
    print(f"Duration: {summary['duration_s']:.2f} s")
    print(f"Sample rate: {summary['sample_rate_hz']:.2f} Hz")
    print(f"Label: {summary['label']}")
    print(f"CSV  : {paths['csv']}")
    print(f"XLSX : {paths['xlsx']}")
    print(f"PNG  : {paths['png']}")
    print(f"LOG  : {paths['log']}")


if __name__ == "__main__":
    main()
